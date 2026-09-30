"""Smoke tests for figures and report builders."""
import io

from docx import Document

from modules.engine import analyze_text, codes_df, risk_df, theme_df
from modules.graphs import dimension_bar, ecosystem_network, risk_radar, sunburst
from modules.parser import extract_text
from modules.reporting import html_document, html_report, pdf_report

TEXT = ("Victims were recruited on Telegram with promises of guaranteed profit. "
        "Funds moved through mule accounts before the police investigation began.")


def _frames():
    a = analyze_text(TEXT)
    return a, theme_df(a), risk_df(a), codes_df(a)


def test_figures_build():
    a, tdf, rdf, _ = _frames()
    for fig in (dimension_bar(tdf), risk_radar(rdf), sunburst(tdf), ecosystem_network()):
        assert fig.to_dict()["data"]


def test_ecosystem_network_edges_all_drawn():
    fig = ecosystem_network()
    # one arrow annotation per directed edge (16) plus one label per node (13)
    arrows = [a for a in fig.layout.annotations if a.showarrow]
    assert len(arrows) == 16 and len(fig.layout.annotations) == 16 + 13


def test_reports_build():
    a, tdf, rdf, cdf = _frames()
    html = html_report(tdf, rdf, cdf, a)
    pdf = pdf_report(tdf, rdf, cdf, a)
    doc = html_document(html)
    assert doc.startswith("<!DOCTYPE html>") and "<meta charset='utf-8'>" in doc and html in doc
    pdf_bytes = pdf if isinstance(pdf, (bytes, bytearray)) else pdf.getvalue()
    assert pdf_bytes[:4] == b"%PDF"


class _Upload(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name


def test_parser_txt_and_docx():
    assert extract_text(_Upload(TEXT.encode(), "a.txt")) == TEXT
    doc = Document()
    doc.add_paragraph("Paragraph text")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Cell A", "Cell B"
    buf = io.BytesIO(); doc.save(buf)
    out = extract_text(_Upload(buf.getvalue(), "b.docx"))
    assert "Paragraph text" in out and "Cell A | Cell B" in out
    assert extract_text(None) == ""


def test_sunburst_handles_more_dimensions_than_palette(monkeypatch):
    import modules.engine as eng
    import modules.graphs as gr
    extra = dict(eng.THEME_KEYWORDS, **{"Extra Dimension": ["extra"]})
    monkeypatch.setattr(gr, "THEME_KEYWORDS", extra)
    monkeypatch.setattr(eng, "THEME_KEYWORDS", extra)
    a = eng.analyze_text(TEXT + " An extra term appears here.")
    assert gr.sunburst(eng.theme_df(a)).to_dict()["data"]
